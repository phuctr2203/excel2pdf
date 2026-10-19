"""Generate the sample workbooks used by the tests and the visual checks.

    python tests/make_samples.py [output_dir]
"""

from __future__ import annotations

import datetime as dt
import io
import math
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import (AreaChart, BarChart, BubbleChart, DoughnutChart, LineChart, PieChart,
                            RadarChart, Reference, ScatterChart, Series)
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, GradientFill, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties


def formatting_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Formatting"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 22

    ws["A1"] = "Quarterly Report"
    ws["A1"].font = Font(size=18, bold=True, color="1F4E79")
    ws.merge_cells("A1:E1")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    ws["A1"].fill = PatternFill("solid", fgColor="DDEBF7")

    headers = ["Item", "Amount", "Date", "Percent", "Note"]
    thin = Side(style="thin", color="808080")
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4472C4")
        c.alignment = Alignment(horizontal="center")
        c.border = Border(top=thin, bottom=Side(style="medium"), left=thin, right=thin)
    rows = [
        ("Revenue", 1234567.891, dt.date(2026, 1, 15), 0.256, "Strong quarter"),
        ("Costs", -456789.12, dt.date(2026, 2, 28), -0.042, "Includes one-off restructuring expenses that overflow"),
        ("Tiền lương (Vietnamese)", 98765.4, dt.datetime(2026, 3, 31, 14, 30), 0.5, "Unicode: Xin chào thế giới"),
        ("Zero", 0, dt.date(2026, 4, 1), 0, None),
        ("Large", 123456789012345, dt.date(2026, 5, 5), 1.5, "中文 日本語 한국어"),
    ]
    for r, row in enumerate(rows, start=4):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.cell(row=r, column=2).number_format = '_("$"* #,##0.00_);_("$"* \\(#,##0.00\\);_("$"* "-"??_);_(@_)'
        ws.cell(row=r, column=3).number_format = "dd-mmm-yyyy"
        ws.cell(row=r, column=4).number_format = "0.0%;[Red]-0.0%"
    ws["B5"].font = Font(color="C00000")

    # number formats gallery
    ws["A10"] = "Number formats"
    ws["A10"].font = Font(bold=True, underline="single")
    samples = [
        ("General", 1234.5678), ("0.00", 3.14159), ("#,##0", 9876543), ("0.00E+00", 0.000123),
        ("# ?/?", 5.25), ("[h]:mm:ss", 1.75), ("dddd, mmmm d, yyyy", dt.date(2026, 10, 19)),
        ("h:mm AM/PM", dt.time(15, 45)), ('"Qty: "#,##0', 42), ("00000", 123), ("[Blue]#,##0;[Red]-#,##0", -77),
    ]
    for i, (fmt, val) in enumerate(samples, start=11):
        ws.cell(row=i, column=1, value=fmt).font = Font(name="Consolas", size=10)
        c = ws.cell(row=i, column=2, value=val)
        c.number_format = fmt

    # alignment and wrapping
    ws["D10"] = "Alignment"
    ws["D10"].font = Font(bold=True)
    ws["D11"] = "Wrapped text that should flow onto several lines inside the cell"
    ws["D11"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("D11:E13")
    ws["D11"].border = Border(left=Side("thick", "FF0000"), right=Side("thick", "FF0000"),
                              top=Side("thick", "FF0000"), bottom=Side("thick", "FF0000"))
    for col in "DE":
        for r in (11, 12, 13):
            ws[f"{col}{r}"].border = Border(left=Side("thick", "FF0000") if col == "D" else None,
                                            right=Side("thick", "FF0000") if col == "E" else None,
                                            top=Side("thick", "FF0000") if r == 11 else None,
                                            bottom=Side("thick", "FF0000") if r == 13 else None)
    ws["D15"] = "Rotated"
    ws["D15"].alignment = Alignment(text_rotation=45, horizontal="center", vertical="center")
    ws.row_dimensions[15].height = 50
    ws["E15"] = "Vertical"
    ws["E15"].alignment = Alignment(text_rotation=90, horizontal="center", vertical="center")
    ws["D17"] = "Centered"
    ws["D17"].alignment = Alignment(horizontal="center", vertical="center")
    ws["E17"] = "Right + indent"
    ws["E17"].alignment = Alignment(horizontal="right", indent=1)
    ws["D18"] = "Shrink to fit this long label"
    ws["D18"].alignment = Alignment(shrink_to_fit=True)

    rich = CellRichText(["Rich ", TextBlock(InlineFont(b=True, color="C00000"), "bold red"), " and ",
                         TextBlock(InlineFont(i=True, sz=14, color="2E75B6"), "italic blue")])
    ws["A23"] = rich
    ws["A24"] = "Strikethrough"
    ws["A24"].font = Font(strike=True)
    ws["A25"] = "Double underline"
    ws["A25"].font = Font(underline="double")
    ws["A26"] = "H"
    ws["A26"].value = CellRichText(["H", TextBlock(InlineFont(vertAlign="subscript"), "2"), "O and x",
                                    TextBlock(InlineFont(vertAlign="superscript"), "2")])
    ws["B23"] = "Visit example.com"
    ws["B23"].hyperlink = "https://example.com"
    ws["B23"].font = Font(color="0563C1", underline="single")

    # borders gallery
    styles = ["hair", "thin", "dotted", "dashed", "dashDot", "dashDotDot", "medium", "mediumDashed",
              "mediumDashDot", "slantDashDot", "thick", "double"]
    ws["A28"] = "Border styles"
    ws["A28"].font = Font(bold=True)
    for i, st in enumerate(styles):
        cell = ws.cell(row=29 + i, column=1, value=st)
        ws.cell(row=29 + i, column=2).border = Border(bottom=Side(style=st, color="000000"))

    # fills
    ws["D28"] = "Fills"
    ws["D28"].font = Font(bold=True)
    for i, pat in enumerate(["solid", "gray125", "lightGray", "darkGray", "lightGrid"]):
        c = ws.cell(row=29 + i, column=4, value=pat)
        c.fill = PatternFill(pat, fgColor="70AD47", bgColor="FFFFFF")
    ws["E29"] = "Gradient"
    ws["E29"].fill = GradientFill(stop=("FFFFFF", "4472C4"), degree=0)
    ws["E30"] = "Theme tint"
    from openpyxl.styles.colors import Color
    ws["E30"].fill = PatternFill("solid", fgColor=Color(theme=5, tint=0.6))
    ws["E31"] = "Diagonal"
    ws["E31"].border = Border(diagonal=Side("thin", "FF0000"), diagonalUp=True, diagonalDown=True)

    # conditional formatting
    ws2 = wb.create_sheet("Conditional")
    ws2["A1"] = "Score"
    ws2["B1"] = "Heat"
    ws2["C1"] = "Bar"
    ws2["D1"] = "Band"
    for i in range(2, 17):
        ws2.cell(row=i, column=1, value=(i * 37) % 100)
        ws2.cell(row=i, column=2, value=(i * 13) % 50)
        ws2.cell(row=i, column=3, value=i * 5)
        ws2.cell(row=i, column=4, value=f"row {i}")
    ws2.conditional_formatting.add("A2:A16", CellIsRule(operator="greaterThan", formula=["70"],
                                                        fill=PatternFill(bgColor="C6EFCE"), font=Font(color="006100")))
    ws2.conditional_formatting.add("A2:A16", CellIsRule(operator="lessThan", formula=["30"],
                                                        fill=PatternFill(bgColor="FFC7CE"), font=Font(color="9C0006")))
    ws2.conditional_formatting.add("B2:B16", ColorScaleRule(start_type="min", start_color="F8696B",
                                                            mid_type="percentile", mid_value=50, mid_color="FFEB84",
                                                            end_type="max", end_color="63BE7B"))
    ws2.conditional_formatting.add("C2:C16", DataBarRule(start_type="min", end_type="max", color="638EC6"))
    ws2.conditional_formatting.add("D2:D16", FormulaRule(formula=["MOD(ROW(),2)=0"], fill=PatternFill(bgColor="EDEDED")))

    # long table across pages with print titles, header/footer
    ws3 = wb.create_sheet("Long table")
    ws3.append(["ID", "Name", "Region", "Q1", "Q2", "Q3", "Q4", "Total", "Comment"])
    for c in ws3[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFF2CC")
        c.border = Border(bottom=Side("medium"))
    regions = ["North", "South", "East", "West"]
    for i in range(1, 121):
        q = [round(1000 + 500 * math.sin(i + k), 2) for k in range(4)]
        ws3.append([i, f"Customer {i:03d}", regions[i % 4], *q, None, "ok" if i % 7 else "check"])
        ws3.cell(row=i + 1, column=8, value=sum(q))
        for col in range(4, 9):
            ws3.cell(row=i + 1, column=col).number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHI", (6, 16, 10, 11, 11, 11, 11, 12, 30)):
        ws3.column_dimensions[col].width = w
    ws3.print_title_rows = "1:1"
    ws3.oddHeader.center.text = "&A"
    ws3.oddFooter.right.text = "Page &P of &N"
    ws3.oddFooter.left.text = "&F"
    ws3.page_setup.paperSize = ws3.PAPERSIZE_A4
    ws3.page_setup.orientation = "landscape"
    ws3.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws3.page_setup.fitToWidth = 1
    ws3.page_setup.fitToHeight = 0
    ws3.print_options.gridLines = True

    ws.oddFooter.center.text = "Page &P of &N"
    wb.save(path)


def _png(w: int, h: int) -> bytes:
    from PIL import Image, ImageDraw

    img = Image.new("RGBA", (w, h), (255, 255, 255, 0))
    d = ImageDraw.Draw(img)
    for i in range(0, w, 8):
        d.line([(i, 0), (w - i, h)], fill=(68, 114, 196, 255), width=2)
    d.ellipse([w * 0.25, h * 0.25, w * 0.75, h * 0.75], fill=(237, 125, 49, 220))
    d.text((8, 8), "PNG", fill=(0, 0, 0, 255))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def charts_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Charts"
    data = [
        ["Month", "North", "South", "East"],
        ["Jan", 120, 90, 60], ["Feb", 135, 85, 72], ["Mar", 150, 110, 80],
        ["Apr", 145, 120, 95], ["May", 170, 130, 105], ["Jun", 190, 125, 118],
    ]
    for row in data:
        ws.append(row)
    for r in range(2, 8):
        for c in range(2, 5):
            ws.cell(row=r, column=c).number_format = "#,##0"
    cats = Reference(ws, min_col=1, min_row=2, max_row=7)
    vals = Reference(ws, min_col=2, min_row=1, max_col=4, max_row=7)

    def place(chart, anchor, title, w=12, h=7):
        chart.title = title
        chart.width, chart.height = w, h
        if hasattr(chart, "x_axis"):
            chart.x_axis.delete = False
            chart.y_axis.delete = False
        ws.add_chart(chart, anchor)

    bar = BarChart()
    bar.add_data(vals, titles_from_data=True)
    bar.set_categories(cats)
    bar.y_axis.majorGridlines = bar.y_axis.majorGridlines
    bar.y_axis.title = "Units"
    place(bar, "F1", "Clustered column")

    sbar = BarChart()
    sbar.type = "bar"
    sbar.grouping = "stacked"
    sbar.overlap = 100
    sbar.add_data(vals, titles_from_data=True)
    sbar.set_categories(cats)
    place(sbar, "F16", "Stacked bar")

    line = LineChart()
    line.add_data(vals, titles_from_data=True)
    line.set_categories(cats)
    for s in line.series:
        s.marker.symbol = "circle"
        s.smooth = False
    line.dataLabels = None
    place(line, "N1", "Line with markers")

    area = AreaChart()
    area.grouping = "stacked"
    area.add_data(vals, titles_from_data=True)
    area.set_categories(cats)
    place(area, "N16", "Stacked area")

    pie = PieChart()
    pie.add_data(Reference(ws, min_col=2, min_row=1, max_row=7), titles_from_data=True)
    pie.set_categories(cats)
    pie.dataLabels = DataLabelList(showPercent=True, showVal=False, showCatName=False,
                                   showSerName=False, showLegendKey=False)
    place(pie, "F31", "Pie (North)")

    dough = DoughnutChart()
    dough.add_data(Reference(ws, min_col=3, min_row=1, max_row=7), titles_from_data=True)
    dough.set_categories(cats)
    place(dough, "N31", "Doughnut (South)")

    # scatter
    ws["A20"] = "x"
    ws["B20"] = "y"
    for i in range(1, 11):
        ws.cell(row=20 + i, column=1, value=i)
        ws.cell(row=20 + i, column=2, value=round(i * i / 3 + 2, 2))
    sc = ScatterChart()
    xs = Reference(ws, min_col=1, min_row=21, max_row=30)
    ys = Reference(ws, min_col=2, min_row=20, max_row=30)
    s = Series(ys, xs, title_from_data=True)
    s.marker.symbol = "diamond"
    sc.series.append(s)
    sc.style = 13
    place(sc, "F46", "Scatter")

    radar = RadarChart()
    radar.type = "marker"
    radar.add_data(vals, titles_from_data=True)
    radar.set_categories(cats)
    place(radar, "N46", "Radar")

    bub = BubbleChart()
    ws["D20"] = "size"
    for i in range(1, 11):
        ws.cell(row=20 + i, column=4, value=i * 3)
    bs = Series(values=Reference(ws, min_col=2, min_row=21, max_row=30),
                xvalues=Reference(ws, min_col=1, min_row=21, max_row=30),
                zvalues=Reference(ws, min_col=4, min_row=21, max_row=30), title="Bubbles")
    bub.series.append(bs)
    place(bub, "F61", "Bubble")

    # bar with data labels + line on secondary axis (combo)
    combo = BarChart()
    combo.add_data(Reference(ws, min_col=2, min_row=1, max_row=7), titles_from_data=True)
    combo.set_categories(cats)
    combo.dataLabels = DataLabelList(showVal=True, showPercent=False, showCatName=False,
                                     showSerName=False, showLegendKey=False)
    ln = LineChart()
    ln.add_data(Reference(ws, min_col=4, min_row=1, max_row=7), titles_from_data=True)
    ln.y_axis.axId = 200
    ln.y_axis.title = "East"
    ln.y_axis.crosses = "max"
    ln.y_axis.delete = False
    combo += ln
    place(combo, "N61", "Combo: column + line (secondary axis)")

    img = XLImage(io.BytesIO(_png(160, 100)))
    ws.add_image(img, "A35")

    ws.page_setup.orientation = "landscape"
    wb.save(path)


def example_workbook(path: Path) -> None:
    """The small report shown in the README: a formatted table plus two charts."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales 2026"
    ws.sheet_view.showGridLines = True
    widths = {"A": 14, "B": 12, "C": 12, "D": 12, "E": 12, "F": 13}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws["A1"] = "Sales Report 2026"
    ws["A1"].font = Font(size=20, bold=True, color="1F4E79")
    ws["A2"] = "Revenue by region and quarter (USD)"
    ws["A2"].font = Font(italic=True, color="595959")

    header = ["Region", "Q1", "Q2", "Q3", "Q4", "Total"]
    thin = Side(style="thin", color="BFBFBF")
    for i, h in enumerate(header, start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.row_dimensions[4].height = 20
    data = [
        ("North", 42150, 45800, 51200, 58900),
        ("South", 38400, 36100, 40750, 44300),
        ("East", 29800, 33500, 35900, 41200),
        ("West", 25600, 27900, 31400, 36800),
    ]
    for r, (region, *qs) in enumerate(data, start=5):
        ws.cell(row=r, column=1, value=region)
        for k, q in enumerate(qs, start=2):
            ws.cell(row=r, column=k, value=q)
        ws.cell(row=r, column=6, value=sum(qs)).font = Font(bold=True)
        for k in range(1, 7):
            cell = ws.cell(row=r, column=k)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if k > 1:
                cell.number_format = "$#,##0"
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="DDEBF7")
    total_row = 9
    ws.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
    for k in range(2, 7):
        col_sum = sum(ws.cell(row=r, column=k).value for r in range(5, 9))
        c = ws.cell(row=total_row, column=k, value=col_sum)
        c.number_format = "$#,##0"
        c.font = Font(bold=True)
    for k in range(1, 7):
        ws.cell(row=total_row, column=k).border = Border(top=Side(style="thin"), bottom=Side(style="double"))
    ws["A11"] = "Growth Q1 → Q4"
    ws["A11"].font = Font(bold=True)
    for r, (region, *qs) in enumerate(data, start=12):
        ws.cell(row=r, column=1, value=region)
        g = ws.cell(row=r, column=2, value=qs[-1] / qs[0] - 1)
        g.number_format = "0.0%"
    ws.conditional_formatting.add("B12:B15", DataBarRule(start_type="num", start_value=0, end_type="max", color="70AD47"))

    bar = BarChart()
    bar.title = "Revenue by quarter"
    bar.add_data(Reference(ws, min_col=2, min_row=4, max_col=5, max_row=8), titles_from_data=True, from_rows=False)
    bar.set_categories(Reference(ws, min_col=1, min_row=5, max_row=8))
    bar.y_axis.numFmt = "$#,##0"
    bar.y_axis.majorGridlines = bar.y_axis.majorGridlines
    bar.x_axis.delete = False
    bar.y_axis.delete = False
    bar.width, bar.height = 15, 7.5
    bar.legend.position = "b"
    ws.add_chart(bar, "H1")

    pie = PieChart()
    pie.title = "Share of annual revenue"
    pie.add_data(Reference(ws, min_col=6, min_row=4, max_row=8), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=5, max_row=8))
    pie.dataLabels = DataLabelList(showPercent=True, showVal=False, showCatName=False,
                                   showSerName=False, showLegendKey=False)
    pie.width, pie.height = 15, 7.5
    ws.add_chart(pie, "H17")

    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.oddFooter.center.text = "Page &P of &N"
    wb.save(path)


def main() -> None:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[2] / "samples"
    out.mkdir(parents=True, exist_ok=True)
    example_workbook(out / "sales-report.xlsx")
    formatting_workbook(out / "formatting.xlsx")
    charts_workbook(out / "charts.xlsx")
    print(f"samples written to {out}")


if __name__ == "__main__":
    main()
