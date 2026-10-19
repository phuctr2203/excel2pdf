"""Public API: convert(source, output=None, **options)."""

from __future__ import annotations

import io
import os
import re
import tempfile
import warnings
import zipfile
from dataclasses import dataclass, field, fields
from pathlib import Path
from typing import BinaryIO, Dict, List, Optional, Sequence, Tuple, Union

import openpyxl
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils import range_boundaries
from reportlab.pdfgen.canvas import Canvas

from . import capture
from .chart import ChartReader
from .colors import ColorResolver
from .drawing import DrawingReader
from .fonts import FontManager
from .package import Package
from .pagination import PageSetup, paper_from_option, setup_from_sheet
from .render import RenderContext, SheetRenderer
from .sheet import StyleCache
from .text import FontSpec, TextEngine
from .theme import Theme

Source = Union[str, os.PathLike, bytes, bytearray, BinaryIO]
Output = Union[str, os.PathLike, BinaryIO, None]

ENGINES = ("native", "capture", "excel", "libreoffice")


@dataclass
class ConvertOptions:
    """Options for :func:`convert`. Page-setup options override what the workbook says."""

    sheets: Optional[Sequence[Union[str, int]]] = None  # names or 0-based indexes; default: all visible
    engine: str = "native"  # native | capture | excel | libreoffice
    paper_size: Union[str, int, Tuple[float, float], None] = None  # "A4", "Letter", Excel code, (w, h) pt
    orientation: Optional[str] = None  # "portrait" | "landscape"
    margins: Union[float, Tuple[float, float, float, float], None] = None  # inches: all or (l, t, r, b)
    scale: Optional[int] = None  # percent, 10-400
    fit_to_width: Optional[int] = None  # pages wide (0 = any)
    fit_to_height: Optional[int] = None  # pages tall (0 = any)
    fit_to_page: bool = False  # shortcut for one page per print area
    gridlines: Optional[bool] = None
    center_horizontally: Optional[bool] = None
    center_vertically: Optional[bool] = None
    ignore_print_area: bool = False
    include_hidden_sheets: bool = False
    images: bool = True
    charts: bool = True
    shapes: bool = True
    conditional_formatting: bool = True
    headers_footers: bool = True
    hyperlinks: bool = True
    fonts: Optional[Dict[str, Union[str, Dict[str, str]]]] = None  # family -> .ttf path
    font_dirs: Optional[Sequence[str]] = None
    use_system_fonts: bool = True
    title: Optional[str] = None  # PDF metadata title
    timeout: float = 180.0  # capture engines only


_FONT_MANAGERS: Dict[tuple, FontManager] = {}


def _font_manager(opts: ConvertOptions) -> FontManager:
    key = (
        tuple(sorted((k, str(v)) for k, v in (opts.fonts or {}).items())),
        tuple(opts.font_dirs or ()),
        opts.use_system_fonts,
    )
    fm = _FONT_MANAGERS.get(key)
    if fm is None:
        fm = FontManager(opts.fonts, opts.font_dirs, opts.use_system_fonts)
        _FONT_MANAGERS[key] = fm
    return fm


def _read_source(source: Source) -> Tuple[bytes, str]:
    if isinstance(source, (bytes, bytearray)):
        return bytes(source), "workbook.xlsx"
    if hasattr(source, "read"):
        data = source.read()  # type: ignore[union-attr]
        name = getattr(source, "name", "workbook.xlsx")
        return data, str(name)
    path = Path(source)  # type: ignore[arg-type]
    return path.read_bytes(), str(path)


def convert(source: Source, output: Output = None, options: Optional[ConvertOptions] = None, **kwargs) -> Optional[bytes]:
    """Convert an Excel workbook to PDF.

    ``source`` is a path, bytes, or a binary file object. ``output`` is a path or a
    binary file object; when omitted the PDF bytes are returned. Keyword arguments
    are the fields of :class:`ConvertOptions`.
    """
    opts = options or ConvertOptions()
    if kwargs:
        valid = {f.name for f in fields(ConvertOptions)}
        unknown = set(kwargs) - valid
        if unknown:
            raise TypeError(f"unknown option(s): {', '.join(sorted(unknown))}")
        opts = ConvertOptions(**{**opts.__dict__, **kwargs})
    if opts.engine not in ENGINES:
        raise ValueError(f"engine must be one of {ENGINES}")

    data, filename = _read_source(source)
    if opts.engine == "native":
        pdf = _convert_native(data, filename, opts)
    else:
        pdf = _convert_capture(data, filename, opts)

    if output is None:
        return pdf
    if hasattr(output, "write"):
        output.write(pdf)  # type: ignore[union-attr]
    else:
        Path(output).write_bytes(pdf)  # type: ignore[arg-type]
    return None


def _convert_capture(data: bytes, filename: str, opts: ConvertOptions) -> bytes:
    suffix = Path(filename).suffix or ".xlsx"
    with tempfile.TemporaryDirectory(prefix="excel2pdf-") as tmp:
        src = Path(tmp) / ("workbook" + suffix)
        dst = Path(tmp) / "workbook.pdf"
        src.write_bytes(data)
        capture.convert(src, dst, opts.engine, opts.timeout)
        return dst.read_bytes()


# --------------------------------------------------------------------------

def _make_resolver(wb):
    ref_re = re.compile(r"^(?:'((?:[^']|'')+)'|([^!]+))!(.+)$")

    def cells_of(sheet_name: str, rng: str):
        if sheet_name.startswith("[") or sheet_name not in wb.sheetnames:
            return None
        ws = wb[sheet_name]
        if not hasattr(ws, "_cells"):
            return None
        rng = rng.replace("$", "")
        if ":" not in rng:
            rng = f"{rng}:{rng}"
        try:
            c0, r0, c1, r1 = range_boundaries(rng)
        except ValueError:
            # a sheet-scoped defined name
            dn = ws.defined_names.get(rng) if hasattr(ws, "defined_names") else None
            return resolve(dn.attr_text) if dn is not None else None
        out = []
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                cell = ws._cells.get((r, c))
                if cell is None:
                    out.append((None, None))
                    continue
                v = cell.value
                if isinstance(v, CellRichText):
                    v = str(v)
                out.append((v, cell.number_format))
        return out

    def resolve(ref: str):
        ref = ref.strip()
        if ref.startswith("(") and ref.endswith(")"):
            ref = ref[1:-1]
        parts = _split_refs(ref)
        if len(parts) > 1:
            out = []
            for p in parts:
                vals = resolve(p)
                if vals is None:
                    return None
                out.extend(vals)
            return out
        m = ref_re.match(ref)
        if not m:
            dn = wb.defined_names.get(ref)
            if dn is not None and dn.attr_text != ref:
                return resolve(dn.attr_text)
            return None
        sheet = (m.group(1) or m.group(2)).replace("''", "'")
        return cells_of(sheet, m.group(3))

    return resolve


def _split_refs(text: str) -> List[str]:
    parts, cur, quoted = [], [], False
    for ch in text:
        if ch == "'":
            quoted = not quoted
        if ch == "," and not quoted:
            parts.append("".join(cur))
            cur = []
        else:
            cur.append(ch)
    parts.append("".join(cur))
    return [p for p in parts if p.strip()]


def _apply_overrides(setup: PageSetup, opts: ConvertOptions) -> PageSetup:
    paper = paper_from_option(opts.paper_size)
    if paper:
        setup.paper = paper
    if opts.orientation:
        o = opts.orientation.lower()
        if o not in ("portrait", "landscape"):
            raise ValueError("orientation must be 'portrait' or 'landscape'")
        setup.landscape = o == "landscape"
    if opts.margins is not None:
        m = opts.margins
        if isinstance(m, (int, float)):
            m = (m, m, m, m)
        setup.margins = tuple(float(v) * 72 for v in m)  # type: ignore[assignment]
    if opts.scale is not None:
        setup.scale = max(10, min(400, int(opts.scale))) / 100.0
        setup.fit = False
    if opts.fit_to_page:
        setup.fit, setup.fit_width, setup.fit_height = True, 1, 1
    if opts.fit_to_width is not None or opts.fit_to_height is not None:
        setup.fit = True
        setup.fit_width = opts.fit_to_width if opts.fit_to_width is not None else 0
        setup.fit_height = opts.fit_to_height if opts.fit_to_height is not None else 0
    if opts.gridlines is not None:
        setup.gridlines = opts.gridlines
    if opts.center_horizontally is not None:
        setup.center_h = opts.center_horizontally
    if opts.center_vertically is not None:
        setup.center_v = opts.center_vertically
    return setup


def _convert_native(data: bytes, filename: str, opts: ConvertOptions) -> bytes:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ValueError(
            "The native engine reads .xlsx/.xlsm files. For legacy .xls files use engine='capture'."
        ) from exc
    pkg = Package(zf)
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", module="openpyxl")
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, rich_text=True, keep_links=False)

    theme = Theme.from_xml(pkg.theme_xml())
    indexed = None
    custom = getattr(wb, "_colors", None)
    if custom:
        indexed = [str(c) for c in custom]
    colors = ColorResolver(theme, indexed)
    fonts = _font_manager(opts)
    text = TextEngine(fonts)
    styles = StyleCache(wb, colors, FontSpec(family=theme.minor_font))
    if wb._fonts:
        styles.default_font = styles.font_spec(wb._fonts[0], FontSpec(family=theme.minor_font))
    date1904 = bool(getattr(wb, "epoch", None) and wb.epoch.year == 1904)
    ctx = RenderContext(
        wb, pkg, theme, colors, text, styles, opts,
        ChartReader(colors, _make_resolver(wb), date1904), DrawingReader(pkg, colors), filename,
    )

    infos = pkg.sheets()
    chosen = []
    for i, info in enumerate(infos):
        if info.kind not in ("worksheet", "chartsheet"):
            continue
        if opts.sheets is not None:
            if not any((isinstance(s, int) and s == i) or (isinstance(s, str) and s == info.name) for s in opts.sheets):
                continue
        elif info.state != "visible" and not opts.include_hidden_sheets:
            continue
        if info.name not in wb.sheetnames:
            continue
        chosen.append(info)
    if opts.sheets is not None:
        missing = [s for s in opts.sheets if not any(
            (isinstance(s, int) and s == i) or (isinstance(s, str) and s == info.name) for i, info in enumerate(infos))]
        if missing:
            raise KeyError(f"sheet(s) not found: {missing}. Available: {[i.name for i in infos]}")

    plans = []
    for info in chosen:
        ws = wb[info.name]
        if info.kind == "chartsheet":
            setup = PageSetup(landscape=True)
            setup = _apply_overrides(setup, opts)
            renderer = SheetRenderer(ctx, ws, info.part, setup, kind="chartsheet")
        else:
            setup = _apply_overrides(setup_from_sheet(ws), opts)
            renderer = SheetRenderer(ctx, ws, info.part, setup)
        pages = renderer.plan()
        if pages:
            plans.append((renderer, pages))

    buf = io.BytesIO()
    canvas = Canvas(buf, pageCompression=1)
    canvas.setCreator("excel2pdf")
    canvas.setProducer("excel2pdf (MIT) with ReportLab")
    title = opts.title or (wb.properties.title if wb.properties and wb.properties.title else Path(filename).stem)
    canvas.setTitle(title)
    if wb.properties and wb.properties.creator:
        canvas.setAuthor(wb.properties.creator)

    total = sum(len(p) for _, p in plans)
    if total == 0:
        warnings.warn("excel2pdf: the workbook has nothing to print; writing a blank page")
        canvas.setPageSize((612, 792))
        canvas.showPage()
    page_no = 0
    for renderer, pages in plans:
        first = renderer.setup.first_page_number
        for idx, page in enumerate(pages):
            page_no += 1
            number = (first + idx) if first else page_no
            renderer.draw_page(canvas, page, number, total, idx)
            canvas.showPage()
    canvas.save()
    return buf.getvalue()
