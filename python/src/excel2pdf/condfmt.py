"""Conditional formatting: evaluate rules per cell and return the adjusted style.

Supported: cellIs, containsText/notContainsText/beginsWith/endsWith, blanks and
errors, top10, aboveAverage, duplicate/unique values, colorScale, dataBar and
expression rules written with a small formula subset (cell references,
arithmetic, comparisons, ROW, COLUMN, MOD, AND, OR, NOT, ISBLANK, ISNUMBER,
ISTEXT, LEN, ABS, INT, LEFT, RIGHT, UPPER, LOWER, TRIM, IF).
"""

from __future__ import annotations

import math
import re
import statistics
from dataclasses import replace
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl.utils import column_index_from_string

from .colors import RGB, ColorResolver, blend
from .sheet import BorderSide, CellStyle, StyleCache
from .text import FontSpec

Bounds = Tuple[int, int, int, int]  # min_col, min_row, max_col, max_row


class _Rule:
    def __init__(self, rule, bounds: List[Bounds]) -> None:
        self.rule = rule
        self.bounds = bounds
        self.type = rule.type
        self.priority = rule.priority or 0
        self.stop = bool(rule.stopIfTrue)
        self._values: Optional[List[float]] = None
        self._all: Optional[List[object]] = None

    def covers(self, r: int, c: int) -> bool:
        for c0, r0, c1, r1 in self.bounds:
            if r0 <= r <= r1 and c0 <= c <= c1:
                return True
        return False

    def top_left(self, r: int, c: int) -> Tuple[int, int]:
        for c0, r0, c1, r1 in self.bounds:
            if r0 <= r <= r1 and c0 <= c <= c1:
                return r0, c0
        b = self.bounds[0]
        return b[1], b[0]


class ConditionalFormats:
    def __init__(self, ws, colors: ColorResolver, styles: StyleCache) -> None:
        self.ws = ws
        self.colors = colors
        self.styles = styles
        self.rules: List[_Rule] = []
        try:
            for cf in ws.conditional_formatting:
                bounds = [(rg.min_col, rg.min_row, rg.max_col, rg.max_row) for rg in cf.sqref.ranges]
                for rule in cf.rules:
                    self.rules.append(_Rule(rule, bounds))
        except Exception:
            self.rules = []
        self.rules.sort(key=lambda r: r.priority)
        self._cache: Dict[tuple, CellStyle] = {}

    # ------------------------------------------------------------------
    def _range_values(self, rule: _Rule) -> List[object]:
        if rule._all is None:
            vals: List[object] = []
            cells = self.ws._cells
            for c0, r0, c1, r1 in rule.bounds:
                r1 = min(r1, self.ws.max_row)
                c1 = min(c1, self.ws.max_column)
                if (r1 - r0 + 1) * (c1 - c0 + 1) > 2_000_000:
                    continue
                for r in range(r0, r1 + 1):
                    for c in range(c0, c1 + 1):
                        cell = cells.get((r, c))
                        if cell is not None and cell.value is not None:
                            vals.append(cell.value)
            rule._all = vals
            rule._values = [float(v) for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
        return rule._all

    def _numbers(self, rule: _Rule) -> List[float]:
        self._range_values(rule)
        return rule._values or []

    # ------------------------------------------------------------------
    def apply(self, r: int, c: int, cell, st: Optional[CellStyle]) -> Optional[CellStyle]:
        if not self.rules or cell is None:
            return st
        matched = []
        for rule in self.rules:
            if not rule.covers(r, c):
                continue
            try:
                result = self._evaluate(rule, r, c, cell)
            except Exception:
                result = None
            if result:
                matched.append((rule, result))
                if rule.stop:
                    break
        if not matched:
            return st
        base = st or self.styles.get(cell)
        out = replace(base)
        font_set = fill_set = False
        for rule, result in matched:
            kind = rule.type
            if kind == "colorScale" and isinstance(result, tuple):
                if not fill_set:
                    out.fill = result
                    out.gradient = None
                    fill_set = True
                continue
            if kind == "dataBar" and isinstance(result, tuple):
                if out.data_bar is None:
                    out.data_bar = result
                continue
            dxf = rule.rule.dxf
            if dxf is None:
                continue
            if dxf.font is not None and not font_set:
                out.font = self._font(dxf.font, out.font)
                font_set = True
            if dxf.fill is not None and not fill_set:
                fill = self._fill(dxf.fill)
                if fill is not None:
                    out.fill = fill
                    out.gradient = None
                    fill_set = True
            if dxf.border is not None:
                for side in ("left", "right", "top", "bottom"):
                    s = getattr(dxf.border, side, None)
                    if s is not None and s.style:
                        setattr(out, side, BorderSide(s.style, self.colors.xl(s.color, (0, 0, 0)) or (0, 0, 0)))
        return out

    def _font(self, f, base: FontSpec) -> FontSpec:
        color = self.colors.xl(f.color, None)
        return FontSpec(
            family=base.family, size=base.size,
            bold=bool(f.b) if f.b is not None else base.bold,
            italic=bool(f.i) if f.i is not None else base.italic,
            underline=f.u if f.u else base.underline,
            strike=bool(f.strike) if f.strike is not None else base.strike,
            color=color if color is not None else base.color,
            vert=base.vert,
        )

    def _fill(self, fill) -> Optional[RGB]:
        ptype = getattr(fill, "fill_type", None) or getattr(fill, "patternType", None)
        bg = self.colors.xl(getattr(fill, "bgColor", None), None)
        fg = self.colors.xl(getattr(fill, "fgColor", None), None)
        if ptype in (None, "solid"):
            # in differential formats the solid colour lives in bgColor
            return bg if bg is not None else fg
        if ptype == "none":
            return None
        return blend(bg or (1, 1, 1), fg or (0, 0, 0), 0.5)

    # ------------------------------------------------------------------
    def _evaluate(self, rule: _Rule, r: int, c: int, cell):
        rl = rule.rule
        v = cell.value
        t = rule.type
        if t == "cellIs":
            if v is None:
                v = 0
            args = [self._formula(f, rule, r, c) for f in (rl.formula or [])]
            op = rl.operator
            return _compare(op, v, args)
        if t in ("containsText", "notContainsText", "beginsWith", "endsWith"):
            text = (rl.text or "").lower()
            s = "" if v is None else str(v).lower()
            if t == "containsText":
                return text in s
            if t == "notContainsText":
                return text not in s
            if t == "beginsWith":
                return s.startswith(text)
            return s.endswith(text)
        if t == "containsBlanks":
            return v is None or str(v).strip() == ""
        if t == "notContainsBlanks":
            return not (v is None or str(v).strip() == "")
        if t == "containsErrors":
            return cell.data_type == "e"
        if t == "notContainsErrors":
            return cell.data_type != "e"
        if t in ("top10", "aboveAverage", "colorScale", "dataBar"):
            if not isinstance(v, (int, float)) or isinstance(v, bool):
                return None
            nums = self._numbers(rule)
            if not nums:
                return None
            v = float(v)
            if t == "top10":
                rank = int(rl.rank or 10)
                n = max(1, int(math.floor(len(nums) * rank / 100.0))) if rl.percent else rank
                ordered = sorted(nums, reverse=not rl.bottom)
                cutoff = ordered[min(n, len(ordered)) - 1]
                return v <= cutoff if rl.bottom else v >= cutoff
            if t == "aboveAverage":
                avg = statistics.fmean(nums)
                above = rl.aboveAverage is None or bool(rl.aboveAverage)
                if rl.stdDev:
                    sd = statistics.pstdev(nums) * float(rl.stdDev)
                    return v >= avg + sd if above else v <= avg - sd
                if rl.equalAverage:
                    return v >= avg if above else v <= avg
                return v > avg if above else v < avg
            if t == "colorScale":
                return self._color_scale(rl.colorScale, nums, v, rule, r, c)
            return self._data_bar(rl.dataBar, nums, v, rule, r, c)
        if t in ("duplicateValues", "uniqueValues"):
            if v is None:
                return False
            vals = self._range_values(rule)
            key = str(v).lower() if isinstance(v, str) else v
            count = sum(1 for x in vals if (str(x).lower() if isinstance(x, str) else x) == key)
            return count > 1 if t == "duplicateValues" else count == 1
        if t == "expression":
            if not rl.formula:
                return None
            return _truthy(self._formula(rl.formula[0], rule, r, c))
        return None

    def _cfvo_value(self, cfvo, nums: List[float], rule, r, c) -> float:
        kind = cfvo.type
        lo, hi = min(nums), max(nums)
        if kind == "min":
            return lo
        if kind == "max":
            return hi
        val = cfvo.val
        if kind == "num":
            x = self._formula(str(val), rule, r, c) if val is not None else 0
            return float(x) if isinstance(x, (int, float)) else lo
        if kind == "percent":
            return lo + (hi - lo) * float(val or 0) / 100.0
        if kind == "percentile":
            p = float(val or 0) / 100.0
            s = sorted(nums)
            k = (len(s) - 1) * p
            f = math.floor(k)
            cc = min(f + 1, len(s) - 1)
            return s[f] + (s[cc] - s[f]) * (k - f)
        if kind == "formula":
            x = self._formula(str(val), rule, r, c)
            return float(x) if isinstance(x, (int, float)) else lo
        return lo

    def _color_scale(self, cs, nums, v, rule, r, c):
        if cs is None or not cs.cfvo or not cs.color:
            return None
        points = [self._cfvo_value(cf, nums, rule, r, c) for cf in cs.cfvo]
        colors = [self.colors.xl(col, (1, 1, 1)) or (1, 1, 1) for col in cs.color]
        if v <= points[0]:
            return colors[0]
        if v >= points[-1]:
            return colors[-1]
        for i in range(len(points) - 1):
            a, b = points[i], points[i + 1]
            if a <= v <= b:
                t = (v - a) / (b - a) if b != a else 0.0
                return blend(colors[i], colors[i + 1], t)
        return colors[-1]

    def _data_bar(self, db, nums, v, rule, r, c):
        if db is None or not db.cfvo:
            return None
        lo = self._cfvo_value(db.cfvo[0], nums, rule, r, c)
        hi = self._cfvo_value(db.cfvo[-1], nums, rule, r, c)
        frac = 0.0 if hi == lo else (v - min(lo, 0.0 if lo > 0 else lo)) / (hi - min(lo, 0.0 if lo > 0 else lo))
        frac = max(0.1, min(1.0, frac))
        color = self.colors.xl(db.color, (0.39, 0.56, 0.86)) or (0.39, 0.56, 0.86)
        return (frac, blend(color, (1, 1, 1), 0.25))

    # ------------------------------------------------------------------
    def _formula(self, text: str, rule: _Rule, r: int, c: int):
        top, left = rule.top_left(r, c)
        return _Formula(text, self.ws, r - top, c - left, (r, c)).evaluate()


def _truthy(v) -> bool:
    if isinstance(v, str):
        return v.upper() == "TRUE"
    return bool(v)


def _compare(op: Optional[str], v, args: Sequence) -> bool:
    def num(x):
        if isinstance(x, bool):
            return float(x)
        if isinstance(x, (int, float)):
            return float(x)
        return x

    v = num(v)
    a = num(args[0]) if args else None
    b = num(args[1]) if len(args) > 1 else None

    def cmp(x, y, f):
        try:
            if isinstance(x, str) and isinstance(y, str):
                return f(x.lower(), y.lower())
            if isinstance(x, str) != isinstance(y, str):
                # Excel: any text sorts after any number
                return f(1 if isinstance(x, str) else 0, 1 if isinstance(y, str) else 0)
            return f(x, y)
        except TypeError:
            return False

    if op == "between":
        lo, hi = (a, b) if cmp(a, b, lambda x, y: x <= y) else (b, a)
        return cmp(v, lo, lambda x, y: x >= y) and cmp(v, hi, lambda x, y: x <= y)
    if op == "notBetween":
        lo, hi = (a, b) if cmp(a, b, lambda x, y: x <= y) else (b, a)
        return not (cmp(v, lo, lambda x, y: x >= y) and cmp(v, hi, lambda x, y: x <= y))
    fns = {
        "equal": lambda x, y: x == y, "notEqual": lambda x, y: x != y,
        "greaterThan": lambda x, y: x > y, "lessThan": lambda x, y: x < y,
        "greaterThanOrEqual": lambda x, y: x >= y, "lessThanOrEqual": lambda x, y: x <= y,
    }
    f = fns.get(op or "")
    return cmp(v, a, f) if f else False


# --------------------------------------------------------------------------
# A tiny formula evaluator for conditional-format expressions
# --------------------------------------------------------------------------

_TOKEN = re.compile(
    r"\s*(?:"
    r"(?P<num>\d+(?:\.\d*)?(?:[eE][-+]?\d+)?|\.\d+)"
    r'|(?P<str>"(?:[^"]|"")*")'
    r"|(?P<ref>(?:(?:'[^']+'|[A-Za-z0-9_]+)!)?\$?[A-Za-z]{1,3}\$?\d+)"
    r"|(?P<func>[A-Za-z][A-Za-z0-9_.]*)\s*\("
    r"|(?P<bool>TRUE|FALSE)\b"
    r"|(?P<op><>|<=|>=|[-+*/^&=<>(),%])"
    r")"
)


class _Formula:
    def __init__(self, text: str, ws, dr: int, dc: int, here: Tuple[int, int] = (1, 1)) -> None:
        text = text.strip()
        if text.startswith("="):
            text = text[1:]
        self.tokens = self._tokenize(text)
        self.i = 0
        self.ws = ws
        self.dr = dr
        self.dc = dc
        self.here = here

    @staticmethod
    def _tokenize(text: str):
        out = []
        pos = 0
        while pos < len(text):
            m = _TOKEN.match(text, pos)
            if not m or m.end() == pos:
                raise ValueError(f"cannot parse {text!r}")
            kind = m.lastgroup
            val = m.group(kind)
            out.append((kind, val))
            pos = m.end()
        return out

    def peek(self):
        return self.tokens[self.i] if self.i < len(self.tokens) else (None, None)

    def take(self):
        tok = self.peek()
        self.i += 1
        return tok

    def evaluate(self):
        v = self.comparison()
        return v

    def comparison(self):
        left = self.concat()
        while self.peek()[0] == "op" and self.peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self.take()[1]
            right = self.concat()
            name = {"=": "equal", "<>": "notEqual", "<": "lessThan", ">": "greaterThan",
                    "<=": "lessThanOrEqual", ">=": "greaterThanOrEqual"}[op]
            left = _compare(name, left if left is not None else 0, [right if right is not None else 0])
        return left

    def concat(self):
        left = self.additive()
        while self.peek() == ("op", "&"):
            self.take()
            left = _text(left) + _text(self.additive())
        return left

    def additive(self):
        left = self.term()
        while self.peek()[0] == "op" and self.peek()[1] in "+-":
            op = self.take()[1]
            right = self.term()
            left = _num(left) + _num(right) if op == "+" else _num(left) - _num(right)
        return left

    def term(self):
        left = self.power()
        while self.peek()[0] == "op" and self.peek()[1] in "*/":
            op = self.take()[1]
            right = self.power()
            left = _num(left) * _num(right) if op == "*" else _num(left) / _num(right)
        return left

    def power(self):
        left = self.unary()
        while self.peek() == ("op", "^"):
            self.take()
            left = _num(left) ** _num(self.unary())
        return left

    def unary(self):
        if self.peek() == ("op", "-"):
            self.take()
            return -_num(self.unary())
        if self.peek() == ("op", "+"):
            self.take()
            return self.unary()
        v = self.primary()
        if self.peek() == ("op", "%"):
            self.take()
            v = _num(v) / 100.0
        return v

    def primary(self):
        kind, val = self.take()
        if kind == "num":
            return float(val)
        if kind == "str":
            return val[1:-1].replace('""', '"')
        if kind == "bool":
            return val.upper() == "TRUE"
        if kind == "ref":
            return self._ref(val)
        if kind == "func":
            name = val.upper()
            args = []
            if self.peek() != ("op", ")"):
                args.append(self.comparison())
                while self.peek() == ("op", ","):
                    self.take()
                    args.append(self.comparison())
            self.take()  # ')'
            return self._call(name, args)
        if (kind, val) == ("op", "("):
            v = self.comparison()
            self.take()
            return v
        raise ValueError(f"unexpected token {val!r}")

    def _ref(self, text: str):
        sheet = None
        if "!" in text:
            sheet, text = text.rsplit("!", 1)
            sheet = sheet.strip("'")
        m = re.match(r"(\$?)([A-Za-z]{1,3})(\$?)(\d+)", text)
        col_abs, col, row_abs, row = m.groups()
        c = column_index_from_string(col.upper()) + (0 if col_abs else self.dc)
        r = int(row) + (0 if row_abs else self.dr)
        ws = self.ws.parent[sheet] if sheet and sheet in self.ws.parent.sheetnames else self.ws
        cell = ws._cells.get((r, c))
        self._last_ref = (r, c)
        return cell.value if cell is not None else None

    def _call(self, name: str, args):
        if name == "ROW":
            return float(self._last_ref[0]) if args and hasattr(self, "_last_ref") else float(self._here()[0])
        if name == "COLUMN":
            return float(self._last_ref[1]) if args and hasattr(self, "_last_ref") else float(self._here()[1])
        if name == "MOD":
            a, b = _num(args[0]), _num(args[1])
            return a - b * math.floor(a / b)
        if name == "AND":
            return all(_truthy(a) for a in args)
        if name == "OR":
            return any(_truthy(a) for a in args)
        if name == "NOT":
            return not _truthy(args[0])
        if name == "IF":
            return args[1] if _truthy(args[0]) else (args[2] if len(args) > 2 else False)
        if name == "ISBLANK":
            return args[0] is None or args[0] == ""
        if name == "ISNUMBER":
            return isinstance(args[0], (int, float)) and not isinstance(args[0], bool)
        if name == "ISTEXT":
            return isinstance(args[0], str)
        if name == "LEN":
            return float(len(_text(args[0])))
        if name == "ABS":
            return abs(_num(args[0]))
        if name == "INT":
            return float(math.floor(_num(args[0])))
        if name == "LEFT":
            return _text(args[0])[: int(_num(args[1])) if len(args) > 1 else 1]
        if name == "RIGHT":
            n = int(_num(args[1])) if len(args) > 1 else 1
            return _text(args[0])[-n:] if n else ""
        if name == "UPPER":
            return _text(args[0]).upper()
        if name == "LOWER":
            return _text(args[0]).lower()
        if name == "TRIM":
            return " ".join(_text(args[0]).split())
        raise ValueError(f"unsupported function {name}")

    def _here(self) -> Tuple[int, int]:
        return self.here


def _num(v) -> float:
    if v is None:
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v))
    except ValueError:
        raise ValueError("#VALUE!")


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)
